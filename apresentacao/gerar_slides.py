"""Gera apresentacao/slides.pptx a partir das figuras em resultados/."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

R = Path(__file__).resolve().parent.parent / "resultados"
FUNDO, TXT, DEST, CINZA = RGBColor(0x0B, 0x10, 0x20), RGBColor(0xF2, 0xF2, 0xF2), RGBColor(0xFF, 0xC8, 0x57), RGBColor(0x9A, 0xA4, 0xB8)

p = Presentation(); p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)


def slide(titulo, quem=None):
    s = p.slides.add_slide(p.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = FUNDO
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.9)).text_frame
    tb.text = titulo or " "; r = tb.paragraphs[0].runs[0]; r.font.size = Pt(32); r.font.bold = True; r.font.color.rgb = TXT
    if quem:
        t = s.shapes.add_textbox(Inches(10.3), Inches(7.0), Inches(2.8), Inches(0.4)).text_frame
        t.text = quem; t.paragraphs[0].runs[0].font.size = Pt(11); t.paragraphs[0].runs[0].font.color.rgb = CINZA
    return s


def texto(s, linhas, x=0.5, y=1.3, w=12.3, h=5.8, tam=20):
    tf = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame; tf.word_wrap = True
    for i, l in enumerate(linhas):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        dest = l.startswith("!"); par.text = l.lstrip("!"); par.space_after = Pt(10)
        for r in par.runs:
            r.font.size = Pt(tam); r.font.color.rgb = DEST if dest else TXT; r.font.bold = dest


def img(s, nome, x, y, w=None, h=None):
    s.shapes.add_picture(str(R / nome), Inches(x), Inches(y),
                         width=Inches(w) if w else None, height=Inches(h) if h else None)


# 1 título
s = slide("")
texto(s, ["!Detecção de estrelas em constelações", "Segmentação + FFT aplicadas a imagens NOIRLab", "",
          "Fernando Fleuri Barbosa · Isabela Xavier Tiosso · Matheus Eduardo Nunhez",
          "Processamento de Imagens e Sinais — Prof. Dr. Vinicius Santos Andrade"], y=2.3, tam=28)

# 2 problema
s = slide("Problema e objetivo", "Isabela")
texto(s, ["Uma foto do céu tem milhares de estrelas fracas, Via Láctea, nebulosas e ruído do sensor.",
          "!Objetivo: isolar as estrelas e destacar as principais, que desenham a constelação.",
          "Aplicação: pré-processamento para star trackers (orientação de satélites), catalogação e apps de astronomia.",
          "Resultado esperado: lista de estrelas (posição, área, brilho) + as 12 mais brilhantes marcadas."], w=6.6, tam=19)
img(s, "10_principais.png", 7.2, 1.4, w=5.8)

# 3 dados
s = slide("Dados de entrada", "Isabela")
img(s, "01_entrada.png", 0.5, 1.2, w=8.2)
texto(s, ["Fonte: NOIRLab — Constellations", "JPEG · RGB 8 bits", "1280 × 1280 px", "Órion e Escorpião",
          "!Histograma: >95% dos pixels são céu escuro; estrelas são a cauda"], x=8.9, w=4.1, tam=17)
img(s, "02_histograma.png", 0.8, 4.4, w=6.0)

# 4 pipeline
s = slide("Pipeline", "Isabela")
etapas = ["RGB → cinza", "FFT 2D", "Passa-faixa\ngaussiano", "IFFT", "Limiar k·σ", "Abertura\n2×2", "Componentes\nconexos", "Ranking\npor fluxo"]
for i, e in enumerate(etapas):
    b = s.shapes.add_shape(1, Inches(0.4 + i * 1.6), Inches(2.6), Inches(1.45), Inches(1.2))
    b.fill.solid(); b.fill.fore_color.rgb = RGBColor(0x1E, 0x2A, 0x48) if i not in (1, 2, 3, 4, 5, 6) else RGBColor(0x2D, 0x4A, 0x8A)
    b.line.color.rgb = DEST; tf = b.text_frame; tf.text = e
    for par in tf.paragraphs:
        for r in par.runs: r.font.size = Pt(14); r.font.color.rgb = TXT
texto(s, ["!Pré-processamento → FFT (frequência) → Segmentação (espacial) → Extração de características → Decisão",
          "Cada etapa recebe a saída da anterior: se a FFT não tira o fundo, o limiar marca a Via Láctea como estrela."],
      y=4.4, tam=18)

# 5 FFT espectro
s = slide("FFT: o que o espectro representa", "Fernando")
img(s, "03_espectro.png", 0.4, 1.2, w=8.6)
texto(s, ["Dado transformado: imagem em cinza (luminância)",
          "!Centro = baixas freq. → fundo, Via Láctea, nebulosas",
          "!Meio = estrelas (PSF 2–8 px)",
          "!Bordas = altas freq. → ruído do sensor, JPEG",
          "Magnitude: quanto de cada frequência · Fase: onde estão as estruturas"], x=9.1, w=4.0, tam=16)
img(s, "04_fase.png", 0.4, 4.3, w=8.6)

# 6 FFT filtragem
s = slide("FFT: filtro passa-faixa", "Fernando")
img(s, "05_mascaras.png", 0.4, 1.15, w=6.4)
img(s, "06_filtragem.png", 0.4, 3.4, w=8.8)
texto(s, ["g = IFFT( FFT(f) · H )", "Gaussiano: sem ringing do filtro ideal",
          "d0 = N/100 tira o fundo", "d1 = N/6 tira o ruído",
          "!Espelhamento nas bordas evita vazamento espectral",
          "Passa-baixa na freq. ≡ Gaussiano espacial (convolução ↔ multiplicação)"], x=9.4, y=1.2, w=3.7, tam=14)

# 7 segmentação
s = slide("Segmentação", "Matheus")
img(s, "08_segmentacao.png", 0.4, 1.15, w=12.5)
texto(s, ["ROI = cada estrela. Limiar = mediana + k·1,4826·MAD (k = 5) → abertura 2×2 → componentes 8-conexos (área ≥ 3 px)",
          "!Otsu supõe 2 classes de tamanho parecido; aqui estrelas são <5% dos pixels → falha. k·σ mede o ruído do fundo."],
      y=5.5, tam=16)

# 8 comparação de métodos
s = slide("Por que essa técnica? Comparação", "Matheus")
img(s, "11_comparacao_metodos.png", 0.4, 1.15, w=12.5)
texto(s, ["K-means agrupa por cor: a nebulosa vira uma classe inteira. Watershed separa estrelas grudadas, mas depende de boa binarização.",
          "!A FFT ataca a causa: tira o fundo antes do limiar."], y=5.6, tam=16)

# 9 resultados
s = slide("Resultados: céu sintético com ground truth", "Fernando")
img(s, "13_f1.png", 0.4, 1.2, w=6.4)
img(s, "14_sensibilidade.png", 6.9, 1.2, w=6.2)
texto(s, ["60 estrelas em posições conhecidas · acerto = detecção a ≤ 4 px",
          "!Gradiente forte: F1 0,17 (original) → 0,99 (FFT + k·σ)",
          "k baixo → ruído vira estrela (precisão cai) · k alto → estrelas fracas somem (recall cai)"], y=4.6, tam=17)

# 10 impacto
s = slide("Impacto de uma segmentação errada", "Isabela")
img(s, "15_impacto.png", 0.4, 1.15, w=8.6)
texto(s, ["Sem FFT, a “estrela” nº 1 é a Via Láctea: 412 mil px.",
          "!O ranking e o desenho da constelação saem errados.",
          "Com FFT: Rigel, Betelgeuse e Bellatrix no topo."], x=9.2, w=3.9, tam=17)

# 11 limitações
s = slide("Limitações e próximos passos", "Matheus")
texto(s, ["Estrelas saturadas (Sírius): núcleo chapado vira baixa frequência → sobra só um anel.",
          "Estrelas sobrepostas em aglomerados → aplicar Watershed depois do passa-faixa.",
          "JPEG de divulgação: o fluxo não é fotometria científica.",
          "!Próximo: identificar a constelação casando triângulos de estrelas com o catálogo Hipparcos.",
          "Removemos o classificador antigo: era treinado com dados aleatórios, sem validade."], tam=20)

# 12 conclusão
s = slide("Conclusão", "Matheus")
texto(s, ["!No domínio da frequência: fundo = baixa · estrela = média · ruído = alta.",
          "O passa-faixa deixa para a segmentação só o que interessa.",
          "F1 de ≈0,2 para ≈0,99 em céu com gradiente; estrelas principais de Órion identificadas.",
          "Código, testes e notebook: github.com/IsaTiosso/processamento-imagens-lab"], y=2.0, tam=24)

p.save(Path(__file__).resolve().parent / "slides.pptx")
print("ok")
